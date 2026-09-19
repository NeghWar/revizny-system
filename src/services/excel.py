"""Export a import dát do/z Excelu (.xlsx).

Export: stylovaný .xlsx súbor (hlavička, mriežka, zmrazený prvý riadok).
Import: načítanie .xlsx, konverzia slovenských dátumov na ISO a očistenie
        float-artefaktov (napr. "17.0" -> "17").
"""

import io

import pandas as pd

from src.config import PERIODA_STLPCE

try:
    from openpyxl.styles import Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    import streamlit as st

    st.error("Chýba knižnica 'openpyxl'. Pridajte ji do requirements.txt")


# Poradie a názvy stĺpcov vo Excel exporte (db_stĺpec -> Excel názov)
STLPCE_PRE_EXCEL: dict[str, str] = {
    "nazov": "Názov stroja",
    "druh_systemu": "Druh systému (VTZ/UTZ)",
    "legislativna_skupina": "Legislatívna skupina",
    "legislativny_druh": "Legislatívny druh",
    "evidencne_cislo": "Evidenčné číslo",
    "cislo_vybavenia": "Číslo vybavenia",
    "umiestnenie": "Umiestnenie",
    "firma": "Firma",
    "voj": "VOJ",
    "nasledujuca_revizia": "Ďalšia Revízia",
    "nasledujuca_revizna_skuska": "Ďalšia Revízna skúška",
    "nasledujuca_podrobna_prehliadka_ok": "Ďalšia Podrobná prehliadka OK",
    "nasledujuca_odborna_prehliadka": "Ďalšia Odborná prehliadka",
    "nasledujuca_odborna_skuska": "Ďalšia Odborná skúška",
    "nasledujuca_uradna_skuska": "Ďalšia Úradná skúška",
    "nasledujuca_geometria": "Ďalšia Geometria dráhy",
    # Periódy jednotlivých kontrol (novinka — doplnené na koniec, aby staré
    # Excel súbory bez týchto stĺpcov fungovali aj naďalej)
    "perioda_reviznej_skusky": "Perióda Revíznej skúšky (roky)",
    "perioda_uradnej_skusky": "Perióda Úradnej skúšky (roky)",
    "interval_odborna_pr": "Interval Odbornéj prehliadky (roky)",
    "perioda_odbornej_skusky": "Perióda Odbornéj skúšky (roky)",
}

# Inverzná mapa pre import: Excel názov -> db stĺpec
MAPOVANIE_NA_DB: dict[str, str] = {excel: db for db, excel in STLPCE_PRE_EXCEL.items()}

# Excel stĺpce, ktoré obsahujú dátumy (naformátované ako DD.MM.YYYY)
STLPCE_S_DATUMAMI: list[str] = [
    excel for db, excel in STLPCE_PRE_EXCEL.items() if db.startswith("nasledujuca_")
]


# ---------------------------------------------------------------------------
# POMOCNÉ KONVERZIE
# ---------------------------------------------------------------------------
def ocisti_float_cislo(hodnota) -> str | None:
    """Odstráni otravné .0 z konca čísel vybavenia a evidencie."""
    if pd.isna(hodnota):
        return None
    hodnota_str = str(hodnota).strip()
    if hodnota_str.endswith(".0"):
        return hodnota_str[:-2]
    return hodnota_str


def konvertuj_na_iso_datum(hodnota) -> str | None:
    """Prevedie slovenský dátum z Excelu na formát YYYY-MM-DD pre Supabase."""
    if pd.isna(hodnota):
        return None
    h_str = str(hodnota).strip()
    if h_str in ["nevykonáva sa", "Nezadané", "Nezadaná", ""]:
        return None
    try:
        parsed_dt = pd.to_datetime(hodnota, dayfirst=True, errors="raise")
        return parsed_dt.strftime("%Y-%m-%d")
    except Exception:
        return None


def konvertuj_periodu(hodnota, db_col: str):
    """Prevedie periódu zo Excelu na db typ (int, alebo float pre interval)."""
    if pd.isna(hodnota):
        return None
    h_str = str(hodnota).strip().replace(",", ".")
    if h_str in ["nevykonáva sa", "Nezadané", "Nezadaná", ""]:
        return None
    try:
        if db_col == "interval_odborna_pr":
            return float(h_str)
        return int(float(h_str))
    except (ValueError, TypeError):
        return None


# ---------------------------------------------------------------------------
# EXPORT
# ---------------------------------------------------------------------------
def export_stroje_excel(stroje: list[dict]) -> bytes:
    """Vytvorí formátovaný Excel súbor z listy strojov a vráti jeho byte obsah."""
    df = pd.DataFrame(stroje)

    # Poistka: ak stĺpec v DataFrame chýba (nie je v DB), vytvoríme ho ako prázdny
    for db_col in STLPCE_PRE_EXCEL.keys():
        if db_col not in df.columns:
            df[db_col] = None

    df_export = df[list(STLPCE_PRE_EXCEL.keys())].rename(columns=STLPCE_PRE_EXCEL)

    # Dátumové stĺpce naformátujeme na DD.MM.YYYY
    for col in df_export.columns:
        if col in STLPCE_S_DATUMAMI:
            df_export[col] = pd.to_datetime(df_export[col], errors="coerce").dt.strftime("%d.%m.%Y")

    df_export = df_export.fillna("Nezadané")

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df_export.to_excel(writer, index=False, sheet_name="Revízie Strojov")
        workbook = writer.book
        worksheet = writer.sheets["Revízie Strojov"]

        hlavicka_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        hlavicka_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        tenka_ciara = Side(border_style="thin", color="D9D9D9")
        mriezka = Border(left=tenka_ciara, right=tenka_ciara, top=tenka_ciara, bottom=tenka_ciara)

        for row in worksheet.iter_rows(min_row=1, max_row=1, min_col=1, max_col=worksheet.max_column):
            for cell in row:
                cell.font = hlavicka_font
                cell.fill = hlavicka_fill
                cell.border = mriezka
        worksheet.freeze_panes = "A2"

        for col in worksheet.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            worksheet.column_dimensions[col_letter].width = max(max_len + 3, 12)
            for cell in col:
                if cell.row > 1:
                    cell.border = mriezka

    return buffer.getvalue()


# ---------------------------------------------------------------------------
# IMPORT
# ---------------------------------------------------------------------------
def parse_import_subor(subor) -> list[dict]:
    """Načíta .xlsx súbor a vráti listu riadkov {nazov, data} pripravenú pre DB.

    Kontrola duplicity a samotný zázpis do databázy robí UI (záložka 4),
    pretože potrebuje aktuálny stav DB.
    """
    df_import = pd.read_excel(subor)

    riadky: list[dict] = []
    for _, row in df_import.iterrows():
        nazov_st = row.get("Názov stroja")
        if pd.isna(nazov_st) or str(nazov_st).strip() == "":
            continue

        stroj_data: dict = {}
        for excel_col, db_col in MAPOVANIE_NA_DB.items():
            if excel_col not in df_import.columns:
                continue
            val = row[excel_col]

            # 1. Dátumové stĺpce -> ISO formát
            if db_col.startswith("nasledujuca_"):
                stroj_data[db_col] = konvertuj_na_iso_datum(val)
            # 2. Čísla vybavenia / evidencie -> očistené od .0
            elif db_col in ["cislo_vybavenia", "evidencne_cislo"]:
                stroj_data[db_col] = ocisti_float_cislo(val)
            # 3. Periódy -> int/float
            elif db_col in PERIODA_STLPCE:
                stroj_data[db_col] = konvertuj_periodu(val, db_col)
            # 4. Ostatné textové stĺpce
            else:
                stroj_data[db_col] = (
                    None
                    if pd.isna(val) or str(val).strip() in ["nevykonáva sa", "Nezadané", "Nezadaná"]
                    else str(val).strip()
                )

        # Prepínač geometrie zapneme podľa zapísaného dátumu
        if "nasledujuca_geometria" in stroj_data and stroj_data["nasledujuca_geometria"] is not None:
            stroj_data["vykonava_sa_geometria"] = True

        riadky.append({"nazov": str(nazov_st).strip(), "data": stroj_data})

    return riadky